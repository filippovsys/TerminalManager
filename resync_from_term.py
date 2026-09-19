# -*- coding: utf-8 -*-
"""
Полная пересинхронизация активных привязок из листа TERM.

В отличие от fix_sync_term.py, этот скрипт УМЕЕТ ДОСОЗДАВАТЬ недостающие:
  * терминалы (по S/N),
  * terminal_ids (по payment_id),
  * и связывать их.
Это лечит последствия неточного импорта.

Что делает:
  1. Читает лист TERM (с учётом объединённых ячеек SAN).
  2. Для каждого M/id -- находит или создаёт мерчанта.
  3. Для каждого S/N -- находит или создаёт терминал.
  4. Для каждого payment_id -- находит или создаёт terminal_ids.
  5. Создаёт/реактивирует активную привязку terminal <-> terminal_id.
  6. Убеждается, что терминал у мерчанта.
  7. Все активные привязки, которых нет в TERM -- закрывает.
  8. Терминалы "у клиента" без активных привязок -- отправляет на склад.

Запуск:
    python resync_from_term.py "СВЕДЕНИЯ IMPORT.xlsx"
"""

import os
import shutil
import sqlite3
import sys
from datetime import datetime

import config


TYPE_MAP = {"telekeçi": "telekeci", "edara": "edara", "bank": "bank"}


def read_term_rows(path):
    """Читает лист TERM. SAN управляет группировкой:
    - при появлении нового SAN -- берём S/N и модель из этой строки;
    - пустые S/N у последующих строк наследуют S/N текущей группы."""
    try:
        import openpyxl
    except ImportError:
        print("Нужна библиотека openpyxl: pip install openpyxl")
        sys.exit(1)

    wb = openpyxl.load_workbook(path, data_only=True)
    if "TERM" not in wb.sheetnames:
        print(f"Нет листа TERM. Есть: {wb.sheetnames}")
        sys.exit(1)
    ws = wb["TERM"]

    rows = []
    current_sn = None
    current_model = None
    for r in ws.iter_rows(min_row=2, values_only=True):
        if len(r) < 15:
            continue
        if r[0] is not None:  # новый SAN
            current_sn = str(r[5]).strip() if r[5] is not None else None
            current_model = r[4]
        payment_id = r[1]
        if payment_id is None:
            continue
        rows.append({
            "payment_id": str(payment_id).strip(),
            "transit_account": str(r[2]).strip() if r[2] is not None else None,
            "m_id": str(r[3]).strip() if r[3] is not None else None,
            "model": r[4] or current_model,
            "sn": (str(r[5]).strip() if r[5] is not None else current_sn),
            "owner_type": r[8],
            "owner_label": r[9],
            "point_label": r[10],
            "address": str(r[11]).strip() if r[11] is not None else None,
            "install_date": r[12],
            "issue_date": r[13],
            "phone": str(r[14]).strip() if r[14] is not None else None,
        })
    return rows


def main():
    if len(sys.argv) < 2:
        print('Использование: python resync_from_term.py "СВЕДЕНИЯ IMPORT.xlsx"')
        return 1
    xlsx_path = sys.argv[1]
    if not os.path.exists(xlsx_path):
        print(f"Файл не найден: {xlsx_path}")
        return 1

    if not os.path.exists(config.DB_PATH):
        print(f"База не найдена: {config.DB_PATH}")
        return 1

    print(f"Excel:  {xlsx_path}")
    print(f"База:   {config.DB_PATH}")
    print()

    print("Читаю лист TERM...")
    rows = read_term_rows(xlsx_path)
    unique_pairs = set((r["payment_id"], r["sn"]) for r in rows)
    print(f"  Строк: {len(rows)}, уникальных пар (ID, S/N): {len(unique_pairs)}")

    # Бэкап
    backup = config.DB_PATH + ".backup-resync-" + datetime.now().strftime("%Y%m%d-%H%M%S")
    shutil.copy2(config.DB_PATH, backup)
    print(f"  Резервная копия: {backup}")
    print()

    conn = sqlite3.connect(config.DB_PATH)
    conn.row_factory = sqlite3.Row

    merchant_cache = {}
    terminal_cache = {}
    terminal_id_cache = {}
    moved_to_merchant = set()
    stats = {
        "merchants_created": 0,
        "terminals_created": 0,
        "terminal_ids_created": 0,
        "bindings_created": 0,
        "bindings_reactivated": 0,
        "skipped_no_sn": 0,
        "skipped_no_mid": 0,
    }

    def get_merchant(m_id, owner_type):
        if m_id is None:
            return None
        if m_id in merchant_cache:
            return merchant_cache[m_id]
        r = conn.execute("SELECT id FROM merchants WHERE m_id = ?", (m_id,)).fetchone()
        if r is not None:
            merchant_cache[m_id] = r["id"]
            return r["id"]
        mtype = "telekeci"
        if owner_type:
            t = str(owner_type).strip().lower()
            if t in TYPE_MAP:
                mtype = TYPE_MAP[t]
        cur = conn.execute(
            "INSERT INTO merchants (m_id, merchant_type, note) VALUES (?, ?, ?)",
            (m_id, mtype, "создан при resync из TERM"),
        )
        merchant_cache[m_id] = cur.lastrowid
        stats["merchants_created"] += 1
        return merchant_cache[m_id]

    def get_terminal(sn, model):
        if sn is None:
            return None
        if sn in terminal_cache:
            return terminal_cache[sn]
        r = conn.execute("SELECT id FROM terminals WHERE serial_number = ?", (sn,)).fetchone()
        if r is not None:
            terminal_cache[sn] = r["id"]
            return r["id"]
        cur = conn.execute(
            "INSERT INTO terminals (serial_number, model, ownership, is_epos) "
            "VALUES (?, ?, 'bank', 0)",
            (sn, model),
        )
        tid = cur.lastrowid
        conn.execute(
            "INSERT INTO terminal_placements (terminal_id, place_type, comment) "
            "VALUES (?, 'warehouse', 'создан при resync из TERM')",
            (tid,),
        )
        terminal_cache[sn] = tid
        stats["terminals_created"] += 1
        return tid

    def get_terminal_id(row, merchant_id):
        pid = row["payment_id"]
        if pid in terminal_id_cache:
            return terminal_id_cache[pid]
        r = conn.execute(
            "SELECT id FROM terminal_ids WHERE payment_id = ? ORDER BY id LIMIT 1",
            (pid,),
        ).fetchone()
        if r is not None:
            terminal_id_cache[pid] = r["id"]
            conn.execute(
                "UPDATE terminal_ids SET "
                "address = COALESCE(address, ?), "
                "phone = COALESCE(phone, ?), "
                "install_date = COALESCE(install_date, ?), "
                "issue_date = COALESCE(issue_date, ?) "
                "WHERE id = ?",
                (row["address"], row["phone"], row["install_date"],
                 row["issue_date"], r["id"]),
            )
            return r["id"]
        cur = conn.execute(
            "INSERT INTO terminal_ids (payment_id, transit_account, merchant_id, "
            "owner_label, point_label, address, phone, install_date, issue_date) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (pid, row["transit_account"], merchant_id,
             row["owner_label"], row["point_label"],
             row["address"], row["phone"],
             row["install_date"], row["issue_date"]),
        )
        terminal_id_cache[pid] = cur.lastrowid
        stats["terminal_ids_created"] += 1
        return cur.lastrowid

    def ensure_binding(terminal_id, terminal_id_ref):
        r = conn.execute(
            "SELECT id, bound_to FROM terminal_id_bindings "
            "WHERE terminal_id = ? AND terminal_id_ref = ? "
            "ORDER BY (bound_to IS NULL) DESC, id DESC LIMIT 1",
            (terminal_id, terminal_id_ref),
        ).fetchone()
        if r is None:
            conn.execute(
                "INSERT INTO terminal_id_bindings (terminal_id, terminal_id_ref) "
                "VALUES (?, ?)",
                (terminal_id, terminal_id_ref),
            )
            stats["bindings_created"] += 1
            return
        if r["bound_to"] is not None:
            conn.execute(
                "UPDATE terminal_id_bindings SET bound_to = NULL WHERE id = ?",
                (r["id"],),
            )
            stats["bindings_reactivated"] += 1

    print("Синхронизирую TERM...")
    for row in rows:
        if row["m_id"] is None:
            stats["skipped_no_mid"] += 1
            continue
        if row["sn"] is None:
            stats["skipped_no_sn"] += 1
            continue
        mid = get_merchant(row["m_id"], row["owner_type"])
        tid = get_terminal(row["sn"], row["model"])
        tid_ref = get_terminal_id(row, mid)
        ensure_binding(tid, tid_ref)

        if tid not in moved_to_merchant:
            st = conn.execute(
                "SELECT current_place, current_merchant_id FROM terminal_current_state "
                "WHERE terminal_id = ?",
                (tid,),
            ).fetchone()
            if st is None or st["current_place"] != "merchant" or st["current_merchant_id"] != mid:
                conn.execute(
                    "INSERT INTO terminal_placements (terminal_id, place_type, merchant_id, comment) "
                    "VALUES (?, 'merchant', ?, 'resync: активен в TERM')",
                    (tid, mid),
                )
            moved_to_merchant.add(tid)

    print(f"  Мерчантов создано:          {stats['merchants_created']}")
    print(f"  Терминалов создано:         {stats['terminals_created']}")
    print(f"  terminal_ids создано:       {stats['terminal_ids_created']}")
    print(f"  Привязок создано:           {stats['bindings_created']}")
    print(f"  Привязок реактивировано:    {stats['bindings_reactivated']}")
    if stats["skipped_no_sn"]:
        print(f"  Пропущено строк без S/N:    {stats['skipped_no_sn']}")
    if stats["skipped_no_mid"]:
        print(f"  Пропущено строк без M/id:   {stats['skipped_no_mid']}")

    # --- 2. Закрыть лишние активные привязки (которых нет в TERM) ---
    print("\nУбираю лишние активные привязки (которых нет в TERM)...")
    term_pairs = set((r["payment_id"], r["sn"]) for r in rows if r["sn"])
    extra = []
    for r in conn.execute("""
        SELECT b.id AS bid, ti.payment_id, t.serial_number
        FROM terminal_id_bindings b
        JOIN terminal_ids ti ON ti.id = b.terminal_id_ref
        JOIN terminals t ON t.id = b.terminal_id
        WHERE b.bound_to IS NULL
    """).fetchall():
        pid = str(r["payment_id"])
        sn = str(r["serial_number"]).strip() if r["serial_number"] is not None else None
        if (pid, sn) not in term_pairs:
            extra.append(r["bid"])
    if extra:
        conn.executemany(
            "UPDATE terminal_id_bindings SET bound_to = datetime('now') WHERE id = ?",
            [(bid,) for bid in extra],
        )
        print(f"  Закрыто лишних привязок: {len(extra)}")
    else:
        print("  Лишних не найдено.")

    # --- 3. Терминалы у клиента без активных ID -- на склад ---
    print("\nПеремещаю на склад терминалы у клиента без активных ID...")
    orphans = conn.execute("""
        SELECT t.id FROM terminals t
        JOIN terminal_current_state s ON s.terminal_id = t.id
        WHERE t.is_epos = 0 AND s.current_place = 'merchant'
          AND t.id NOT IN (
              SELECT terminal_id FROM terminal_id_bindings WHERE bound_to IS NULL
          )
    """).fetchall()
    for o in orphans:
        conn.execute(
            "INSERT INTO terminal_placements (terminal_id, place_type, comment) "
            "VALUES (?, 'warehouse', 'resync: нет активных ID')",
            (o["id"],),
        )
    print(f"  Перемещено: {len(orphans)}")

    conn.commit()

    # --- Итог ---
    print()
    print("=" * 62)
    print("ИТОГОВАЯ СВОДКА")
    print("=" * 62)
    rows2 = conn.execute("""
        SELECT m.merchant_type, COUNT(*) IDs, COUNT(DISTINCT b.terminal_id) Terms
        FROM terminal_id_bindings b
        JOIN terminal_ids ti ON ti.id = b.terminal_id_ref
        JOIN merchants m ON m.id = ti.merchant_id
        JOIN terminals t ON t.id = b.terminal_id
        JOIN terminal_current_state s ON s.terminal_id = b.terminal_id
        WHERE b.bound_to IS NULL AND t.is_epos = 0
          AND s.current_place = 'merchant' AND s.condition != 'written_off'
        GROUP BY m.merchant_type
    """).fetchall()
    tt, ti_sum = 0, 0
    for r in rows2:
        print(f"  {r['merchant_type']}: {r['Terms']} терм / {r['IDs']} ID")
        tt += r["Terms"]; ti_sum += r["IDs"]
    print(f"  ИТОГО: {tt} терм / {ti_sum} ID")

    conn.close()
    print("\nГотово. Запусти программу и проверь вкладку «Активные».")
    return 0


if __name__ == "__main__":
    sys.exit(main())