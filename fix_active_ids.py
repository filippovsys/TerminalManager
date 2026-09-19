# -*- coding: utf-8 -*-
"""
Одноразовая миграция: закрывает активные привязки, которых нет
в листе "TERM" исходного Excel-файла.

Проблема, которую решает: при импорте все строки из листа "ЗАКРЫТЫЕ"
попадали в базу как АКТИВНЫЕ привязки. Правильное поведение:
активны только те (payment_id, S/N), что есть в листе "TERM".

Запуск:
    python fix_active_ids.py "путь\\к\\СВЕДЕНИЯ_о_терминалах.xlsx"

Перед запуском СДЕЛАЙ КОПИЮ базы (terminals.db).
"""

import os
import shutil
import sqlite3
import sys
from datetime import datetime

import config


def read_term_pairs(path):
    """Возвращает set пар (payment_id, serial_number) из листа TERM.
    Учитывает объединённые ячейки столбца S/N (пусто = тянется предыдущий)."""
    try:
        import openpyxl
    except ImportError:
        print("Ошибка: нужна библиотека openpyxl. Установи: pip install openpyxl")
        sys.exit(1)

    wb = openpyxl.load_workbook(path, data_only=True)
    if "TERM" not in wb.sheetnames:
        print(f"Ошибка: в файле нет листа 'TERM'. Есть: {wb.sheetnames}")
        sys.exit(1)

    ws = wb["TERM"]
    pairs = set()
    current_sn = None
    for row in ws.iter_rows(min_row=2, values_only=True):
        if len(row) < 6:
            continue
        if all(c is None for c in row[:6]):
            continue
        payment_id = row[1]  # столбец B -- ID
        sn_cell = row[5]     # столбец F -- S/N
        if sn_cell is not None:
            current_sn = str(sn_cell).strip()
        if payment_id is None:
            continue
        pairs.add((str(payment_id), current_sn))
    return pairs


def main():
    if len(sys.argv) < 2:
        print('Использование: python fix_active_ids.py "путь/к/СВЕДЕНИЯ.xlsx"')
        return 1

    xlsx_path = sys.argv[1]
    if not os.path.exists(xlsx_path):
        print(f"Файл не найден: {xlsx_path}")
        return 1

    if not os.path.exists(config.DB_PATH):
        print(f"База не найдена: {config.DB_PATH}")
        return 1

    print(f"Excel:  {xlsx_path}")
    print(f"База:   {config.DB_PATH}")
    print()

    print("Читаю лист TERM...")
    term_pairs = read_term_pairs(xlsx_path)
    print(f"  Уникальных пар (ID, S/N) в TERM: {len(term_pairs)}")
    print()

    print("Анализирую активные привязки в базе...")
    conn = sqlite3.connect(config.DB_PATH)
    conn.row_factory = sqlite3.Row

    rows = conn.execute("""
        SELECT b.id AS binding_id, ti.payment_id, t.serial_number
        FROM terminal_id_bindings b
        JOIN terminal_ids ti ON ti.id = b.terminal_id_ref
        JOIN terminals t ON t.id = b.terminal_id
        WHERE b.bound_to IS NULL
    """).fetchall()
    print(f"  Активных привязок сейчас: {len(rows)}")

    to_close = []
    for r in rows:
        pid = str(r["payment_id"])
        sn = r["serial_number"]
        if sn is not None:
            sn = str(sn).strip()
        if (pid, sn) not in term_pairs:
            to_close.append((r["binding_id"], pid, sn))

    print(f"  Будет закрыто (нет в TERM): {len(to_close)}")
    print()

    if not to_close:
        print("Нечего делать. Всё уже правильно.")
        conn.close()
        return 0

    print("Примеры того, что будет закрыто (первые 10):")
    for bid, pid, sn in to_close[:10]:
        print(f"  ID={pid}, S/N={sn}")
    print()

    ans = input("Продолжить? (y/n): ").strip().lower()
    if ans != "y":
        print("Отменено.")
        conn.close()
        return 0

    # Резервная копия
    backup_path = config.DB_PATH + ".backup-" + datetime.now().strftime("%Y%m%d-%H%M%S")
    shutil.copy2(config.DB_PATH, backup_path)
    print(f"\nРезервная копия: {backup_path}")

    conn.executemany(
        "UPDATE terminal_id_bindings SET bound_to = datetime('now') WHERE id = ?",
        [(bid,) for bid, _, _ in to_close],
    )
    conn.commit()
    print(f"\nГотово. Закрыто {len(to_close)} привязок.")
    print("Запусти программу -- теперь активных ID ровно столько, сколько в Excel.")

    conn.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())