# -*- coding: utf-8 -*-
"""
Заполняет SN-историю: переносит в базу все ЗАКРЫТЫЕ привязки из листа
"ЗАКРЫТЫЕ" файла Excel.

Что делает для каждой строки из ЗАКРЫТЫЕ:
  * Находит/создаёт мерчанта по M/id.
  * Находит/создаёт терминал по S/N.
  * Находит/создаёт terminal_ids по payment_id.
  * Создаёт ЗАКРЫТУЮ привязку (bound_to = дата из "Состояние" или
    условная).
Если такая привязка уже есть -- пропускает (идемпотентно).

Запуск:
    python fill_closed_history.py "СВЕДЕНИЯ IMPORT.xlsx"
"""

import os
import re
import shutil
import sqlite3
import sys
from datetime import datetime, timedelta

import config


def parse_owner_type(s):
    if not s:
        return None
    s = str(s).strip().lower()
    if "edara" in s:
        return "edara"
    if "bank" in s:
        return "bank"
    if "teleke" in s:
        return "telekeci"
    return None


def parse_closed_date(s):
    """'Закрыт 15.10.2019' -> '2019-10-15 00:00:00'. Иначе None."""
    if not s:
        return None
    m = re.search(r"(\d{1,2})\.(\d{1,2})\.(\d{4})", str(s))
    if m:
        d, mo, y = m.groups()
        return f"{y}-{int(mo):02d}-{int(d):02d} 00:00:00"
    return None


def read_closed_rows(path):
    try:
        import openpyxl
    except ImportError:
        print("Нужна библиотека openpyxl: pip install openpyxl")
        sys.exit(1)

    wb = openpyxl.load_workbook(path, data_only=True)
    if "ЗАКРЫТЫЕ" not in wb.sheetnames:
        print(f"Нет листа ЗАКРЫТЫЕ. Есть: {wb.sheetnames}")
        sys.exit(1)
    ws = wb["ЗАКРЫТЫЕ"]

    rows = []
    current_sn = None
    current_model = None
    for r in ws.iter_rows(min_row=2, values_only=True):
        if len(r) < 6:
            continue
        if r[0] is not None:  # новый SAN
            current_sn = str(r[5]).strip() if r[5] is not None else None
            current_model = r[4]
        payment_id = r[1]
        if payment_id is None:
            continue
        sn = str(r[5]).strip() if r[5] is not None else current_sn
        rows.append({
            "payment_id": str(payment_id).strip(),
            "transit_account": str(r[2]).strip() if r[2] is not None else None,
            "m_id": str(r[3]).strip() if r[3] is not None else None,
            "model": r[4] or current_model,
            "sn": sn,
            "owner_type": r[8] if len(r) > 8 else None,
            "owner_label": r[9] if len(r) > 9 else None,
            "point_label": r[10] if len(r) > 10 else None,
            "address": str(r[11]).strip() if len(r) > 11 and r[11] is not None else None,
            "install_date": r[12] if len(r) > 12 else None,
            "issue_date": r[13] if len(r) > 13 else None,
            "phone": str(r[14]).strip() if len(r) > 14 and r[14] is not None else None,
            "closed_text": r[15] if len(r) > 15 else None,
        })
    return rows


def main():
    if len(sys.argv) < 2:
        print('Использование: python fill_closed_history.py "СВЕДЕНИЯ IMPORT.xlsx"')
        return 1
    xlsx_path = sys.argv[1]
    if not os.path.exists(xlsx_path):
        print(f"Файл не найден: {xlsx_path}")
        return 1

    print("Читаю лист ЗАКРЫТЫЕ...")
    rows = read_closed_rows(xlsx_path)
    print(f"  Строк: {len(rows)}")

    backup = config.DB_PATH + ".backup-closedhist-" + datetime.now().strftime("%Y%m%d-%H%M%S")
    shutil.copy2(config.DB_PATH, backup)
    print(f"  Резервная копия: {backup}")
    print()

    conn = sqlite3.connect(config.DB_PATH)
    conn.row_factory = sqlite3.Row

    merchant_cache = {}
    terminal_cache = {}
    tid_cache = {}
    stats = {"merchants": 0, "terminals": 0, "terminal_ids": 0,
             "bindings_created": 0, "bindings_skipped": 0,
             "no_sn": 0, "no_mid": 0}

    def get_merchant(m_id, owner_type=None):
        if m_id in merchant_cache:
            return merchant_cache[m_id]
        r = conn.execute("SELECT id FROM merchants WHERE m_id = ?", (m_id,)).fetchone()
        if r is not None:
            merchant_cache[m_id] = r["id"]
            return r["id"]
        mtype = parse_owner_type(owner_type) or "telekeci"
        cur = conn.execute(
            "INSERT INTO merchants (m_id, merchant_type, note) VALUES (?, ?, ?)",
            (m_id, mtype, "создан при заполнении SN истории"),
        )
        merchant_cache[m_id] = cur.lastrowid
        stats["merchants"] += 1
        return merchant_cache[m_id]

    def get_terminal(sn, model):
        if sn in terminal_cache:
            return terminal_cache[sn]
        r = conn.execute("SELECT id FROM terminals WHERE serial_number = ?", (sn,)).fetchone()
        if r is not None:
            terminal_cache[sn] = r["id"]
            return r["id"]
        cur = conn.execute(
            "INSERT INTO terminals (serial_number, model, ownership, is_epos) "
            "VALUES (?, ?, 'bank', 0)",
            (sn, model),
        )
        tid = cur.lastrowid
        conn.execute(
            "INSERT INTO terminal_placements (terminal_id, place_type, comment) "
            "VALUES (?, 'warehouse', 'создан для SN истории (из ЗАКРЫТЫЕ)')",
            (tid,),
        )
        terminal_cache[sn] = tid
        stats["terminals"] += 1
        return tid

    def get_terminal_id_ref(row, merchant_id):
        pid = row["payment_id"]
        if pid in tid_cache:
            return tid_cache[pid]
        r = conn.execute(
            "SELECT id FROM terminal_ids WHERE payment_id = ? ORDER BY id LIMIT 1",
            (pid,),
        ).fetchone()
        if r is not None:
            tid_cache[pid] = r["id"]
            return r["id"]
        cur = conn.execute(
            "INSERT INTO terminal_ids (payment_id, transit_account, merchant_id, "
            "owner_label, point_label, address, phone, install_date, issue_date) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (pid, row["transit_account"], merchant_id,
             row["owner_label"], row["point_label"], row["address"], row["phone"],
             row["install_date"], row["issue_date"]),
        )
        tid_cache[pid] = cur.lastrowid
        stats["terminal_ids"] += 1
        return cur.lastrowid

    default_close = (datetime.now() - timedelta(days=365)).strftime("%Y-%m-%d 00:00:00")

    print("Создаю закрытые привязки...")
    for row in rows:
        if row["sn"] is None:
            stats["no_sn"] += 1
            continue
        if row["m_id"] is None:
            stats["no_mid"] += 1
            continue

        mid = get_merchant(row["m_id"], row["owner_type"])
        tid = get_terminal(row["sn"], row["model"])
        tid_ref = get_terminal_id_ref(row, mid)

        # Проверяем, есть ли уже привязка (terminal, tid_ref)
        r = conn.execute(
            "SELECT id FROM terminal_id_bindings "
            "WHERE terminal_id = ? AND terminal_id_ref = ? LIMIT 1",
            (tid, tid_ref),
        ).fetchone()
        if r is not None:
            stats["bindings_skipped"] += 1
            continue

        closed_at = parse_closed_date(row["closed_text"]) or default_close
        bound_from = row["install_date"] or "2010-01-01 00:00:00"
        if isinstance(bound_from, datetime):
            bound_from = bound_from.strftime("%Y-%m-%d %H:%M:%S")

        conn.execute(
            "INSERT INTO terminal_id_bindings "
            "(terminal_id, terminal_id_ref, bound_from, bound_to) "
            "VALUES (?, ?, ?, ?)",
            (tid, tid_ref, bound_from, closed_at),
        )
        stats["bindings_created"] += 1

    conn.commit()
    conn.close()

    print()
    print("=" * 62)
    print("ГОТОВО")
    print("=" * 62)
    print(f"  Мерчантов создано:        {stats['merchants']}")
    print(f"  Терминалов создано:       {stats['terminals']}")
    print(f"  terminal_ids создано:     {stats['terminal_ids']}")
    print(f"  Закрытых привязок:        {stats['bindings_created']}")
    print(f"  Пропущено (уже есть):     {stats['bindings_skipped']}")
    print(f"  Пропущено без S/N:        {stats['no_sn']}")
    print(f"  Пропущено без M/id:       {stats['no_mid']}")
    print()
    print("Запусти программу, открой карточку терминала -- вкладка «SN история»")
    print("должна заполниться историческими привязками.")
    return 0


if __name__ == "__main__":
    sys.exit(main())