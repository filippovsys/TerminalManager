# -*- coding: utf-8 -*-
"""Отчёт по моделям терминалов: сводка + экспорт в Excel."""

import tkinter as tk
from tkinter import ttk, messagebox, filedialog

import database


class ModelsReportWindow(tk.Toplevel):
    def __init__(self, master, user):
        super().__init__(master)
        self.user = user
        self.title("Отчёт по моделям терминалов")
        self.geometry("1180x680")
        self._row_by_iid = {}

        with database.get_connection() as conn:
            self.rows = database.get_models_report(conn)

        self._build()
        self.transient(master)
        self.grab_set()

    def _build(self):
        # --- Сводка сверху ---
        total_terms = sum(r["terminals_total"] for r in self.rows)
        total_active = sum(r["active_ids"] for r in self.rows)
        total_models = len(self.rows)

        stats = ttk.LabelFrame(self, text="Сводка", padding=8)
        stats.pack(fill="x", padx=10, pady=(10, 4))
        ttk.Label(stats, text=f"Моделей: {total_models}",
                  font=("TkDefaultFont", 10, "bold")).pack(side="left", padx=15)
        ttk.Label(stats, text=f"Всего терминалов: {total_terms}",
                  font=("TkDefaultFont", 10, "bold")).pack(side="left", padx=15)
        ttk.Label(stats, text=f"Активных ID: {total_active}",
                  font=("TkDefaultFont", 10, "bold")).pack(side="left", padx=15)

        # --- Таблица ---
        cols = (
            "model",
            "terminals_total", "active_ids",
            "place_merchant", "place_warehouse", "place_repair_shop", "place_none",
            "cond_normal", "cond_in_repair", "cond_written_off",
            "own_bank", "own_client",
            "clients_bank", "clients_edara", "clients_telekeci", "clients_other",
        )
        headings = {
            "model": "Модель",
            "terminals_total": "Терм.",
            "active_ids": "Акт.ID",
            "place_merchant": "У клиента",
            "place_warehouse": "Склад",
            "place_repair_shop": "Мастер.",
            "place_none": "Без места",
            "cond_normal": "Норма",
            "cond_in_repair": "В ремонте",
            "cond_written_off": "Списано",
            "own_bank": "Банк",
            "own_client": "Клиент",
            "clients_bank": "BANK",
            "clients_edara": "EDARA",
            "clients_telekeci": "TELEKEÇI",
            "clients_other": "Прочие",
        }
        widths = {
            "model": 180,
            "terminals_total": 60, "active_ids": 65,
            "place_merchant": 80, "place_warehouse": 70,
            "place_repair_shop": 70, "place_none": 75,
            "cond_normal": 70, "cond_in_repair": 80, "cond_written_off": 75,
            "own_bank": 60, "own_client": 65,
            "clients_bank": 70, "clients_edara": 70,
            "clients_telekeci": 85, "clients_other": 70,
        }
        self.tree = ttk.Treeview(self, columns=cols, show="headings", height=20)
        for c in cols:
            self.tree.heading(c, text=headings[c])
            anchor = "w" if c == "model" else "center"
            self.tree.column(c, width=widths[c], anchor=anchor)
        self.tree.pack(fill="both", expand=True, padx=10, pady=(0, 4))
        self.tree.tag_configure("total", font=("TkDefaultFont", 9, "bold"),
                                 background="#eef3f9")

        self._fill_tree()

        # --- Кнопки ---
        btns = ttk.Frame(self, padding=(10, 0, 10, 10))
        btns.pack(fill="x")
        ttk.Button(btns, text="Экспорт в Excel...", command=self._export).pack(side="left", padx=2)
        ttk.Button(btns, text="Обновить", command=self._reload).pack(side="left", padx=2)
        ttk.Button(btns, text="Закрыть", command=self.destroy).pack(side="right", padx=2)

    def _fill_tree(self):
        self.tree.delete(*self.tree.get_children())
        self._row_by_iid.clear()
        for r in self.rows:
            iid = self.tree.insert("", "end", values=(
                r["model"],
                r["terminals_total"], r["active_ids"],
                r["place_merchant"], r["place_warehouse"],
                r["place_repair_shop"], r["place_none"],
                r["cond_normal"], r["cond_in_repair"], r["cond_written_off"],
                r["own_bank"], r["own_client"],
                r["clients_bank"], r["clients_edara"],
                r["clients_telekeci"], r["clients_other"],
            ))

        # Итоговая строка
        if self.rows:
            total = {k: 0 for k in self.rows[0] if k != "model"}
            for r in self.rows:
                for k, v in r.items():
                    if k != "model" and isinstance(v, int):
                        total[k] += v
            self.tree.insert("", "end", values=(
                "ИТОГО",
                total["terminals_total"], total["active_ids"],
                total["place_merchant"], total["place_warehouse"],
                total["place_repair_shop"], total["place_none"],
                total["cond_normal"], total["cond_in_repair"], total["cond_written_off"],
                total["own_bank"], total["own_client"],
                total["clients_bank"], total["clients_edara"],
                total["clients_telekeci"], total["clients_other"],
            ), tags=("total",))

    def _reload(self):
        with database.get_connection() as conn:
            self.rows = database.get_models_report(conn)
        self._fill_tree()

    def _export(self):
        path = filedialog.asksaveasfilename(
            parent=self,
            title="Экспорт отчёта",
            defaultextension=".xlsx",
            filetypes=[("Excel файлы", "*.xlsx")],
            initialfile="Отчёт_по_моделям.xlsx",
        )
        if not path:
            return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter

            wb = Workbook()
            ws = wb.active
            ws.title = "По моделям"

            header_font = Font(bold=True)
            header_fill = PatternFill("solid", fgColor="D9E2EF")
            total_font = Font(bold=True)
            total_fill = PatternFill("solid", fgColor="EEF3F9")

            # Заголовки
            headers = [
                "Модель", "Терминалов", "Активных ID",
                "У клиента", "Склад", "Мастерская", "Без места",
                "Норма", "В ремонте", "Списано",
                "Банк", "Клиент",
                "BANK", "EDARA", "TELEKEÇI", "Прочие",
            ]
            for i, h in enumerate(headers, start=1):
                c = ws.cell(row=1, column=i, value=h)
                c.font = header_font
                c.fill = header_fill
                c.alignment = Alignment(horizontal="center")

            row = 2
            total = {k: 0 for k in self.rows[0] if k != "model"} if self.rows else {}
            for r in self.rows:
                ws.cell(row=row, column=1, value=r["model"])
                values = [
                    r["terminals_total"], r["active_ids"],
                    r["place_merchant"], r["place_warehouse"],
                    r["place_repair_shop"], r["place_none"],
                    r["cond_normal"], r["cond_in_repair"], r["cond_written_off"],
                    r["own_bank"], r["own_client"],
                    r["clients_bank"], r["clients_edara"],
                    r["clients_telekeci"], r["clients_other"],
                ]
                for i, v in enumerate(values, start=2):
                    ws.cell(row=row, column=i, value=v).alignment = Alignment(horizontal="center")
                for k, v in r.items():
                    if k != "model" and isinstance(v, int):
                        total[k] += v
                row += 1

            # Итоговая строка
            if total:
                ws.cell(row=row, column=1, value="ИТОГО").font = total_font
                total_values = [
                    total["terminals_total"], total["active_ids"],
                    total["place_merchant"], total["place_warehouse"],
                    total["place_repair_shop"], total["place_none"],
                    total["cond_normal"], total["cond_in_repair"], total["cond_written_off"],
                    total["own_bank"], total["own_client"],
                    total["clients_bank"], total["clients_edara"],
                    total["clients_telekeci"], total["clients_other"],
                ]
                for i, v in enumerate(total_values, start=2):
                    c = ws.cell(row=row, column=i, value=v)
                    c.font = total_font
                    c.fill = total_fill
                    c.alignment = Alignment(horizontal="center")

            ws.freeze_panes = "A2"
            for i in range(1, len(headers) + 1):
                ws.column_dimensions[get_column_letter(i)].width = 14
            ws.column_dimensions["A"].width = 22

            wb.save(path)
        except Exception as exc:
            messagebox.showerror("Ошибка экспорта", str(exc), parent=self)
            return
        messagebox.showinfo("Готово", f"Файл сохранён:\n{path}", parent=self)