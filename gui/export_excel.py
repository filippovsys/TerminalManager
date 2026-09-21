# -*- coding: utf-8 -*-
"""Экспорт текущего дерева главного окна в Excel.

Обходит все строки Treeview (включая свёрнутые), сохраняет:
  - заголовки колонок,
  - текст в #0 (для вкладок «Активные» / «E-POS» это № и S/N),
  - значения колонок,
  - жирным + синей заливкой -- строки-группы (тип клиента),
  - жирным + светлой заливкой -- строки терминалов.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


_SHEET_NAMES = {
    "active": "Активные",
    "warehouse": "Склад",
    "written_off": "Списанные",
    "epos": "E-POS",
}

_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="D9E2EF")
_GROUP_FILL = PatternFill("solid", fgColor="D9E2EF")
_TERM_FILL = PatternFill("solid", fgColor="EEF3F9")
_WRAP = Alignment(vertical="top", wrap_text=False)


def _visible_columns(tree):
    display = tree["displaycolumns"]
    if display == "#all" or display == ("#all",):
        return list(tree["columns"])
    return list(display)


def export_tree_to_xlsx(tree, tab_key, path):
    """Экспортирует всё дерево `tree` в xlsx-файл `path`."""
    wb = Workbook()
    ws = wb.active
    ws.title = _SHEET_NAMES.get(tab_key, "Экспорт")

    include_tree_col = tab_key in ("active", "epos")
    visible_cols = _visible_columns(tree)

    # ----- Заголовки -----
    col = 1
    if include_tree_col:
        ws.cell(row=1, column=col, value=tree.heading("#0")["text"] or "№ / S/N / ID")
        col += 1
    for c in visible_cols:
        title = tree.heading(c)["text"] or c
        ws.cell(row=1, column=col, value=title)
        col += 1

    total_cols = col - 1
    for i in range(1, total_cols + 1):
        cell = ws.cell(row=1, column=i)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # ----- Обход дерева -----
    row_idx = [2]

    def walk(parent_iid):
        for iid in tree.get_children(parent_iid):
            item = tree.item(iid)
            values = item.get("values") or []
            text0 = item.get("text") or ""
            tags = item.get("tags") or ()

            col = 1
            if include_tree_col:
                ws.cell(row=row_idx[0], column=col, value=text0)
                col += 1
            for i, c in enumerate(visible_cols):
                v = values[i] if i < len(values) else ""
                ws.cell(row=row_idx[0], column=col + i, value=v)

            # Стиль по уровню
            if "group" in tags:
                for c2 in range(1, total_cols + 1):
                    cc = ws.cell(row=row_idx[0], column=c2)
                    cc.font = _HEADER_FONT
                    cc.fill = _GROUP_FILL
            elif "terminal" in tags:
                for c2 in range(1, total_cols + 1):
                    cc = ws.cell(row=row_idx[0], column=c2)
                    cc.font = Font(bold=True)
                    cc.fill = _TERM_FILL
            else:
                for c2 in range(1, total_cols + 1):
                    ws.cell(row=row_idx[0], column=c2).alignment = _WRAP

            row_idx[0] += 1
            walk(iid)

    walk("")

    # ----- Авто-ширина -----
    for i in range(1, total_cols + 1):
        max_len = 0
        for r in range(1, row_idx[0]):
            v = ws.cell(row=r, column=i).value
            if v is None:
                continue
            max_len = max(max_len, len(str(v)))
        ws.column_dimensions[get_column_letter(i)].width = min(max(max_len + 2, 8), 55)

    # ----- Закрепить первую строку -----
    ws.freeze_panes = "A2"

    wb.save(path)